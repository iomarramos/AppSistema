#!/usr/bin/env python3
"""Convierte los archivos del SGP de planificación teórica, plan real, realizado y requisición (datos/plan_real/origen)
a CSV ordenados y verifica sus fórmulas. No inventa datos: lo que no cuadra se lista en validacion.txt.

Uso: python3 herramientas/convertir_plan_real.py
Requiere openpyxl.
"""
import csv
import datetime as dt
import os
import re
import unicodedata
from collections import defaultdict

import openpyxl

RAIZ = os.path.join(os.path.dirname(__file__), '..', 'datos')
ORIGEN = os.path.join(RAIZ, 'plan_real', 'origen')
SALIDA = os.path.join(RAIZ, 'plan_real')

MENU = os.path.join(ORIGEN, 'MENU_REAL_Y__TEORICO.xlsx')
COSTOS = os.path.join(ORIGEN, 'Costo_Plan._Teorico_-_Plan._Real_-_Realizado_Alimentacion.xlsx')
REQUISICION = os.path.join(ORIGEN, 'REQUISICION.xlsx')

# (archivo, hoja) -> (nivel, servicio). "Hoja7" no tiene nombre: por su estructura (25 componentes, como la cena) es el
# plan real de la cena; repite agosto y no trae octubre, que llegó aparte en MENU_REAL_CENA.xlsx (2026-10-03).
HOJAS_MENU = {
    ('MENU_REAL_Y__TEORICO.xlsx', 'DESAYUNO_TEORICA'): ('TEORICO', 'DESAYUNO'),
    ('MENU_REAL_Y__TEORICO.xlsx', 'ALMUERZO_TEORICA'): ('TEORICO', 'ALMUERZO'),
    ('MENU_REAL_Y__TEORICO.xlsx', 'CENA_TEORICA'): ('TEORICO', 'CENA'),
    ('MENU_REAL_Y__TEORICO.xlsx', 'DESAYUNO_REAL'): ('REAL', 'DESAYUNO'),
    ('MENU_REAL_Y__TEORICO.xlsx', 'ALMUERZO_REAL'): ('REAL', 'ALMUERZO'),
    ('MENU_REAL_Y__TEORICO.xlsx', 'Hoja7'): ('REAL', 'CENA'),
    ('MENU_REAL_CENA.xlsx', 'Hoja2'): ('REAL', 'CENA'),
}

observaciones = []
validacion = []


def norm(texto):
    t = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', t).strip().upper()


def vacio(v):
    return v is None or (isinstance(v, str) and v.strip() == '')


def fecha_de(texto):
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', str(texto))
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None


def num(v):
    return None if vacio(v) else float(v)


def escribir(nombre, encabezado, filas):
    with open(os.path.join(SALIDA, nombre), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(encabezado)
        w.writerows(filas)


def fmt(v, dec=6):
    if v is None:
        return ''
    if isinstance(v, float):
        return f'{v:.{dec}f}'.rstrip('0').rstrip('.') if '.' in f'{v:.{dec}f}' else f'{v:.{dec}f}'
    return str(v)


# ---------------------------------------------------------------- menú teórico y plan real
def leer_menu():
    libros = {}
    platos, dias = [], []
    vistos = {}
    for (archivo, hoja), (nivel, servicio) in HOJAS_MENU.items():
        if archivo not in libros:
            libros[archivo] = openpyxl.load_workbook(os.path.join(ORIGEN, archivo), data_only=True)
        filas = list(libros[archivo][hoja].iter_rows(values_only=True))
        inicios = [i for i, r in enumerate(filas) if r[1] == 'Estructura Servicio']
        for ini in inicios:
            enc = filas[ini]
            fin = next(i for i in range(ini, len(filas)) if filas[i][1] == 'Comensales')
            cols = [j for j, c in enumerate(enc) if fecha_de(c)]
            fechas = [fecha_de(enc[j]) for j in cols]
            mes = (nivel, servicio, fechas[0].strftime('%Y-%m'))
            firma = tuple(str(filas[k][cols[0]]) for k in range(ini + 1, fin))
            if mes in vistos:
                igual = vistos[mes] == firma
                observaciones.append(f'{archivo} / {hoja}: el bloque de la fila {ini + 1} repite {nivel} {servicio} {mes[2]}'
                                     + (' con el mismo contenido (se omite)' if igual else ' con otro contenido (se omite)'))
                continue
            vistos[mes] = firma
            for j, f in zip(cols, fechas):
                # Columnas del día: bandera R en j-1, receta en j y luego los encabezados (N.Rac., Cto.Plato o Por.(%) y Costo, Cod. Receta).
                nombres = {}
                k = j + 1
                while k < len(enc) and not vacio(enc[k]) and not fecha_de(enc[k]):
                    nombres[str(enc[k]).strip()] = k
                    k += 1
                costo_dia = num(filas[ini - 1][j - 1])
                com = filas[fin]
                comensales = num(com[nombres['N.Rac.']])
                comensales2 = num(com[nombres['Por.(%)']]) if 'Por.(%)' in nombres else None
                dias.append([nivel, servicio, f.isoformat(), fmt(costo_dia), fmt(comensales), fmt(comensales2)])
                for orden, i in enumerate(range(ini + 1, fin), 1):
                    r = filas[i]
                    receta = r[j]
                    if vacio(receta):
                        continue
                    cod = str(r[nombres['Cod. Receta']] or '').split('&')[0].strip()
                    costo = num(r[nombres['Cto.Plato']]) if 'Cto.Plato' in nombres else num(r[nombres['Costo']])
                    por = num(r[nombres['Por.(%)']]) if 'Por.(%)' in nombres else None
                    platos.append([nivel, servicio, f.isoformat(), orden, str(r[1]).strip(), cod, str(receta).strip(),
                                   fmt(num(r[nombres['N.Rac.']])), fmt(por, 4), fmt(costo), str(r[j - 1] or '').strip()])
    return platos, dias


# ---------------------------------------------------------------- costo teórico, plan real y realizado
def leer_costos():
    filas = list(openpyxl.load_workbook(COSTOS, data_only=True).active.iter_rows(values_only=True))
    titulo = next(str(r[1]) for r in filas if r[1] and 'Costo Plan' in str(r[1]))
    cab, salida, totales = {}, [], []
    for r in filas:
        if r[0] in ('Contrato', 'Regimen', 'Servicio'):
            cab[r[0]] = str(r[1]).strip()
        elif isinstance(r[0], dt.datetime):
            salida.append([cab['Contrato'], cab['Regimen'], cab['Servicio'], r[0].date().isoformat()] + [fmt(num(v)) for v in r[1:12]])
        elif r[0] in ('Total', 'T. General'):
            totales.append([r[0], cab.get('Regimen', ''), cab.get('Servicio', '')] + [fmt(num(v)) for v in r[1:12]])
    return titulo, salida, totales


# ---------------------------------------------------------------- requisición
def leer_requisicion():
    filas = list(openpyxl.load_workbook(REQUISICION, data_only=True).active.iter_rows(values_only=True))
    salida, preparacion = [], {}
    regimen = servicio = fecha = estructura = None
    receta_cod = receta = raciones = None
    en_prep = False
    for r in filas:
        a, b, c = r[0], r[1], r[2]
        if isinstance(a, str) and a.startswith('Regimen :'):
            m = re.match(r'Regimen : (\d+) (.*?)- Servicio : (\d+) (.*?) - Fecha : (\d{2}/\d{2}/\d{4})', a)
            regimen, servicio, fecha = m.group(2).strip(), m.group(4).strip(), fecha_de(m.group(5)).isoformat()
            en_prep = False
            continue
        if b == 'Preparación':
            en_prep = True
            if c and receta_cod:
                preparacion.setdefault((receta_cod, receta), []).append(str(c).strip())
            continue
        if en_prep and vacio(b) and c and receta_cod:
            preparacion.setdefault((receta_cod, receta), []).append(str(c).strip())
            continue
        if vacio(b):
            continue
        if b == 'Código':
            continue
        if isinstance(b, str) and vacio(c) and not re.match(r'^\d', b):
            estructura, en_prep = b.strip(), False
            continue
        texto = str(c or '')
        m = re.match(r'^(\S.*?) \[N[º°o]\. Rac\. (\d+(?:\.\d+)?)\]', texto)
        if m and vacio(r[3]):
            receta_cod, receta, raciones, en_prep = str(b).strip(), m.group(1).strip(), float(m.group(2)), False
            continue
        if texto.startswith('  ') and receta_cod:
            salida.append([regimen, servicio, fecha, estructura, receta_cod, receta, fmt(raciones), str(b).strip(), texto.strip(),
                           fmt(num(r[3])), fmt(num(r[4])), str(r[5] or '').strip(), fmt(num(r[6])), str(r[7] or '').strip()])
    return salida, preparacion


def enlazar_productos(requisicion):
    catalogo = {}
    with open(os.path.join(RAIZ, 'sgp', 'catalogo_sgp.csv'), encoding='utf-8-sig') as f:
        for r in csv.DictReader(f, delimiter=';'):
            catalogo.setdefault(norm(r['producto_descripcion']), r['producto_codigo'])
    codigos = {}
    for r in requisicion:
        codigos.setdefault(r[7], (r[8], r[11], r[13]))
    filas = []
    for cod, (desc, ub, ud) in sorted(codigos.items(), key=lambda x: x[1][0]):
        filas.append([cod, desc, ub, ud, catalogo.get(norm(desc), '')])
    return filas


def enlazar_recetas(platos, requisicion):
    app = {}
    with open(os.path.join(RAIZ, 'recetas', 'recetas_normalizadas.csv'), encoding='utf-8-sig') as f:
        for r in csv.DictReader(f, delimiter=';'):
            app.setdefault(norm(r['receta_nombre']), r['receta_codigo'])
    sgp = {}
    for p in platos:
        if p[5]:
            sgp.setdefault(p[5], p[6])
    for r in requisicion:
        sgp.setdefault(r[4], r[5])
    return [[cod, nombre, app.get(norm(nombre), '')] for cod, nombre in sorted(sgp.items(), key=lambda x: x[1])]


# ---------------------------------------------------------------- costo piso y techo por factores
def piso_techo(platos, dias):
    """Por nivel, servicio y mes: factor de cada componente = Σ raciones ÷ Σ comensales del mes. Con ese factor,
    piso = Σ factor × costo por ración más barato del componente; techo = Σ factor × el más caro; medio = Σ raciones ×
    costo ÷ Σ comensales. Devuelve también cuántos días quedaron dentro de la banda."""
    filas = []
    claves = sorted({(d[0], d[1], d[2][:7]) for d in dias})
    for nivel, servicio, mes in claves:
        ds = [d for d in dias if (d[0], d[1], d[2][:7]) == (nivel, servicio, mes) and d[4]]
        com = sum(float(d[4]) for d in ds)
        if com <= 0:
            continue
        comp = defaultdict(lambda: [0.0, [], 0.0])
        for p in platos:
            if (p[0], p[1], p[2][:7]) == (nivel, servicio, mes) and p[7] and p[9]:
                c = comp[p[4]]
                c[0] += float(p[7]); c[1].append(float(p[9])); c[2] += float(p[7]) * float(p[9])
        piso = sum(c[0] / com * min(c[1]) for c in comp.values())
        techo = sum(c[0] / com * max(c[1]) for c in comp.values())
        medio = sum(c[2] for c in comp.values()) / com
        costos = [float(d[3]) for d in ds if d[3]]
        dentro = sum(1 for x in costos if piso - 0.005 <= x <= techo + 0.005)
        filas.append([nivel, servicio, mes, len(comp), f'{piso:.2f}', f'{medio:.2f}', f'{techo:.2f}',
                      f'{min(costos):.2f}', f'{max(costos):.2f}', f'{dentro}/{len(costos)}'])
    return filas


# ---------------------------------------------------------------- verificaciones
def verificar(platos, dias, costos, requisicion):
    def ok(nombre, buenos, total, detalle=''):
        validacion.append(f'{"OK " if buenos == total else "REV"} {nombre}: {buenos} de {total}{" — " + detalle if detalle else ""}')

    # 1. Costo minuta día = suma(raciones × costo por ración) / comensales (teórico).
    por_dia = defaultdict(float)
    for p in platos:
        if p[0] == 'TEORICO' and p[7] and p[9]:
            por_dia[(p[1], p[2])] += float(p[7]) * float(p[9])
    buenos = total = 0
    for d in dias:
        if d[0] == 'TEORICO' and d[3] and d[4] and float(d[4]) > 0:
            total += 1
            buenos += abs(por_dia[(d[1], d[2])] / float(d[4]) - float(d[3])) <= 0.02
    ok('Teórico: costo minuta día = Σ(raciones × costo por ración) ÷ comensales (±0,02)', buenos, total)

    # 2. Plan real: porcentaje = raciones ÷ comensales.
    com = {(d[1], d[2]): float(d[4]) for d in dias if d[0] == 'REAL' and d[4]}
    buenos = total = 0
    for p in platos:
        if p[0] == 'REAL' and p[7] and p[8] and com.get((p[1], p[2])):
            total += 1
            buenos += abs(float(p[7]) / com[(p[1], p[2])] - float(p[8])) <= 0.0005
    ok('Plan real: Por.(%) = raciones ÷ comensales (±0,0005)', buenos, total)

    # 3. Plan real: el "Costo" de cada plato es costo de la ración (el costo minuta día lo confirma).
    por_dia = defaultdict(float)
    for p in platos:
        if p[0] == 'REAL' and p[7] and p[9]:
            por_dia[(p[1], p[2])] += float(p[7]) * float(p[9])
    buenos = total = 0
    for d in dias:
        if d[0] == 'REAL' and d[3] and d[4] and float(d[4]) > 0:
            total += 1
            buenos += abs(por_dia[(d[1], d[2])] / float(d[4]) - float(d[3])) <= 0.02
    ok('Plan real: costo minuta día = Σ(raciones × costo por ración) ÷ comensales (±0,02)', buenos, total)

    # 4. Comparativo: desviaciones y costo total.
    b1 = b2 = b3 = total = 0
    for c in costos:
        v = [float(x) if x else None for x in c[4:15]]
        cbt, rt, ctt, cbr, rr, ctr, dpl, cbz, rz, ctz, drz = v
        if None in (cbt, cbr, dpl):
            continue
        total += 1
        b1 += abs((cbr - cbt) - dpl) <= 0.011
        # Sin realizado (día aún no ejecutado: raciones realizadas 0) el SGP muestra desviación 0.
        b2 += cbz is None or drz is None or not rz or not cbz or abs((cbz - cbr) - drz) <= 0.011
        b3 += rt in (None, 0) or abs(ctt / rt - cbt) <= 0.011
    ok('Comparativo: desviación plan = costo bandeja real − teórico', b1, total)
    ok('Comparativo: desviación realizado = costo bandeja realizado − real (días con realizado)', b2, total)
    ok('Comparativo: costo bandeja = costo total ÷ raciones', b3, total)

    # 5. El comparativo de agosto coincide con los menús (costo minuta día).
    menu = {(d[0], d[1], d[2]): float(d[3]) for d in dias if d[3]}
    mapa = {'DESAYUNO NORMAL 1': 'DESAYUNO', 'ALMUERZO NORMAL 1': 'ALMUERZO', 'CENA NORMAL 1': 'CENA'}
    for nivel, col in (('TEORICO', 4), ('REAL', 7)):
        buenos = total = 0
        dif = []
        for c in costos:
            s = mapa.get(c[2])
            if not s or (nivel, s, c[3]) not in menu or not c[col]:
                continue
            total += 1
            if abs(menu[(nivel, s, c[3])] - float(c[col])) <= 0.011:
                buenos += 1
            elif len(dif) < 3:
                dif.append(f'{s} {c[3]}: menú {menu[(nivel, s, c[3])]} vs comparativo {c[col]}')
        ok(f'Comparativo {nivel.lower()} = costo minuta día del menú', buenos, total, '; '.join(dif))

    # 6. Requisición: cantidad despacho = cantidad bruta × raciones.
    buenos = total = 0
    for r in requisicion:
        if r[9] and r[12] and r[6]:
            total += 1
            # La cantidad por ración se muestra con 4 decimales: el error de redondeo llega a 0,00005 × raciones.
            buenos += abs(float(r[9]) * float(r[6]) - float(r[12])) <= 0.00005 * float(r[6]) + 0.0011
    ok('Requisición: despacho = cantidad bruta por ración × raciones (con el redondeo a 4 decimales)', buenos, total)

    # 7. Requisición contra el plan real: mismas raciones por receta y día.
    plan = {}
    for p in platos:
        if p[0] == 'REAL' and p[5] and p[7]:
            plan.setdefault((p[1], p[2], p[5]), set()).add(float(p[7]))
    rec = {(mapa.get(r[1]), r[2], r[4], float(r[6])) for r in requisicion if r[6] and mapa.get(r[1])}
    buenos = total = 0
    distintos = []
    for s, f, cod, rac in sorted(rec):
        if (s, f, cod) in plan:
            total += 1
            if rac in plan[(s, f, cod)]:
                buenos += 1
            elif len(distintos) < 5:
                distintos.append(f'{s} {f} receta {cod}: requisición {rac:g}, plan real {sorted(plan[(s, f, cod)])}')
    sin_plan = len({(s, f) for s, f, _, _ in rec if not any(k[0] == s and k[1] == f for k in plan)})
    ok('Requisición: raciones = las del plan real del mismo servicio, día y receta', buenos, total,
       '; '.join(distintos + ([f'{sin_plan} servicios-día de la requisición no tienen plan real en el archivo'] if sin_plan else [])))


def main():
    platos, dias = leer_menu()
    titulo, costos, totales = leer_costos()
    requisicion, preparacion = leer_requisicion()

    escribir('menu_planificado.csv', ['nivel', 'servicio', 'fecha', 'orden', 'estructura', 'receta_codigo_sgp', 'receta', 'raciones',
                                      'porcentaje', 'costo_racion', 'marca'], platos)
    escribir('menu_dias.csv', ['nivel', 'servicio', 'fecha', 'costo_minuta_dia', 'comensales', 'comensales_2'], dias)
    escribir('comparativo_costos.csv', ['contrato', 'regimen', 'servicio', 'fecha',
                                        'teorico_costo_bandeja', 'teorico_raciones', 'teorico_costo_total',
                                        'real_costo_bandeja', 'real_raciones', 'real_costo_total', 'desviacion_plan',
                                        'realizado_costo_bandeja', 'realizado_raciones', 'realizado_costo_total', 'desviacion_realizado'], costos)
    escribir('comparativo_totales.csv', ['total', 'regimen', 'servicio',
                                         'teorico_costo_bandeja', 'teorico_raciones', 'teorico_costo_total',
                                         'real_costo_bandeja', 'real_raciones', 'real_costo_total', 'desviacion_plan',
                                         'realizado_costo_bandeja', 'realizado_raciones', 'realizado_costo_total', 'desviacion_realizado'], totales)
    escribir('requisicion.csv', ['regimen', 'servicio', 'fecha', 'estructura', 'receta_codigo_sgp', 'receta', 'raciones',
                                 'producto_codigo_sgp', 'producto', 'cantidad_bruta_racion', 'cantidad_bulto', 'unidad_bulto',
                                 'cantidad_despacho', 'unidad_despacho'], requisicion)
    escribir('preparacion_recetas.csv', ['receta_codigo_sgp', 'receta', 'paso'],
             [[k[0], k[1], p] for k, pasos in sorted(preparacion.items()) for p in dict.fromkeys(pasos)])
    productos = enlazar_productos(requisicion)
    escribir('codigos_sgp_productos.csv', ['producto_codigo_sgp', 'descripcion', 'unidad_bulto', 'unidad_despacho', 'producto_codigo_app'], productos)
    recetas = enlazar_recetas(platos, requisicion)
    escribir('codigos_sgp_recetas.csv', ['receta_codigo_sgp', 'receta', 'receta_codigo_app'], recetas)

    banda = piso_techo(platos, dias)
    escribir('costo_piso_techo.csv', ['nivel', 'servicio', 'mes', 'componentes', 'costo_piso', 'costo_medio', 'costo_techo',
                                      'dia_mas_barato', 'dia_mas_caro', 'dias_dentro'], banda)
    verificar(platos, dias, costos, requisicion)
    validacion.append(('OK ' if all(b[9].split('/')[0] == b[9].split('/')[1] for b in banda) else 'REV')
                      + f' Costo piso y techo por factores: los días de cada servicio y mes quedan dentro de su banda '
                        f'({sum(int(b[9].split("/")[0]) for b in banda)} de {sum(int(b[9].split("/")[1]) for b in banda)})')

    resumen = [
        titulo,
        f'Menú: {len(platos)} platos en {len(dias)} días-servicio '
        f'({", ".join(f"{n} {s}: {sum(1 for d in dias if d[0] == n and d[1] == s)}" for n, s in sorted(set((d[0], d[1]) for d in dias)))})',
        f'Comparativo: {len(costos)} filas por día en {len(totales) - 1} servicios',
        f'Requisición: {len(requisicion)} líneas de producto, {len({(r[1], r[2]) for r in requisicion})} servicios-día, '
        f'{len({r[4] for r in requisicion})} recetas',
        f'Productos de la requisición enlazados al catálogo por descripción: {sum(1 for p in productos if p[4])} de {len(productos)}',
        f'Recetas SGP enlazadas a las recetas cargadas por nombre: {sum(1 for r in recetas if r[2])} de {len(recetas)}',
        '', 'Verificaciones:'] + validacion + ['', 'Observaciones:'] + (observaciones or ['(ninguna)'])
    with open(os.path.join(SALIDA, 'validacion.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(resumen) + '\n')
    print('\n'.join(resumen))


if __name__ == '__main__':
    main()
